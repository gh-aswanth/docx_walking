"""Selective accept / reject.

``accept_all`` proves the markup is semantically right; these tests prove it
stays right when a reviewer resolves *one* revision and leaves the rest as live
tracked changes -- which is the state Word hands back after a partial review.
"""

import docx
import pytest
from conftest import DATE, roundtrip, texts

from docx_redline import Redliner, accept_file, make_selector, reject_file
from docx_redline.editing import review
from docx_redline.oxml.ns import qn


def reopen(rl: Redliner) -> Redliner:
    """Save and reload, so every assertion runs against real serialised markup."""
    return Redliner(roundtrip(rl), track_changes=False)


@pytest.fixture
def two_edits(rl):
    """Two independent edits: a replacement, and a whole new paragraph."""
    rl.replace_text("thirty (30) days", "forty-five (45) days")
    rl.insert_paragraph_after(rl.find_paragraph(contains="4. Termination"), "5. Notices.")
    return rl


def ids_for(rl: Redliner, *kinds: str) -> list[str]:
    return [r.id for r in rl.summary().revisions if r.kind in kinds]


def position(body: list[str], needle: str) -> int:
    """Index of the one paragraph containing ``needle`` -- to assert on ordering."""
    return body.index(next(line for line in body if needle in line))


# ---------------------------------------------------------------------------
# ids are reported
# ---------------------------------------------------------------------------
def test_summary_reports_ids(two_edits):
    summary = two_edits.summary()
    assert summary.ids
    assert len(set(summary.ids)) == len(summary.ids)  # the id pool never collides
    for rev in summary.revisions:
        assert summary.by_id(rev.id) is rev
    assert summary.by_id("no-such-id") is None


def test_summary_format_shows_ids(two_edits):
    text = two_edits.summary().format()
    assert f"#{two_edits.summary().ids[0]} " in text


def test_summary_reports_ids_in_headers(rl_agreement):
    section = rl_agreement.document.sections[0]
    section.header.is_linked_to_previous = False
    section.header.paragraphs[0].text = "Confidential draft"
    rl_agreement.replace_text("Confidential", "Privileged")
    header = [r for r in rl_agreement.summary().revisions if r.location == "hdr"]
    assert header and all(r.id for r in header)


# ---------------------------------------------------------------------------
# selecting by id
# ---------------------------------------------------------------------------
def test_accept_one_id_leaves_the_others_tracked(two_edits, original_text):
    summary = two_edits.summary()
    insert = next(r for r in summary.revisions if r.text == "forty-five (45) days")
    delete = next(r for r in summary.revisions if r.text == "thirty (30) days")

    rl = reopen(two_edits)
    rl.accept(ids=[insert.id, delete.id])

    after = rl.summary()
    assert "forty-five (45) days" in "\n".join(texts(rl.document))
    assert {r.text for r in after.revisions} == {"5. Notices.", ""}  # only the new ¶ is left
    # and the survivor still resolves both ways on its own
    accepted = reopen(rl)
    accepted.accept_all()
    assert "5. Notices." in "\n".join(texts(accepted.document))
    rejected = reopen(rl)
    rejected.reject_all()
    assert "5. Notices." not in "\n".join(texts(rejected.document))


def test_reject_one_id_leaves_the_others_tracked(two_edits, original_text):
    rl = reopen(two_edits)
    rl.reject(ids=ids_for(two_edits, "insert", "paragraph-mark-insert"))

    assert "5. Notices." not in "\n".join(texts(rl.document))
    assert {r.kind for r in rl.summary().revisions} == {"delete"}


def test_a_paragraph_mark_is_its_own_revision(two_edits):
    """Inserting a paragraph is two revisions -- the runs, and the ¶ mark."""
    runs = next(r for r in two_edits.summary().revisions if r.text == "5. Notices.")
    rl = reopen(two_edits)
    rl.accept(ids=runs.id)

    assert rl.summary().counts["paragraph-mark-insert"] == 1  # still tracked on its own
    assert "5. Notices." in "\n".join(texts(rl.document))


def test_ids_accept_ints_and_a_bare_value(two_edits):
    rev = two_edits.summary().revisions[0]
    rl = reopen(two_edits)
    rl.accept(ids=int(rev.id))
    assert rev.id not in reopen(rl).summary().ids


def test_unmatched_id_changes_nothing(two_edits):
    before = two_edits.summary()
    rl = reopen(two_edits)
    rl.accept(ids="999999")
    assert len(rl.summary()) == len(before)


@pytest.mark.parametrize("reverse", [False, True], ids=["in-order", "reversed"])
def test_one_at_a_time_equals_all_at_once(two_edits, reverse):
    """Order must not matter: piecemeal review lands where accept_all does."""
    order = two_edits.summary().ids
    piecemeal = reopen(two_edits)
    for rev_id in reversed(order) if reverse else order:
        piecemeal.accept(ids=rev_id)

    wholesale = reopen(two_edits)
    wholesale.accept_all()
    assert texts(piecemeal.document) == texts(wholesale.document)
    assert len(piecemeal.summary()) == 0


@pytest.mark.parametrize("reverse", [False, True], ids=["in-order", "reversed"])
def test_one_at_a_time_equals_all_at_once_on_reject(two_edits, original_text, reverse):
    order = two_edits.summary().ids
    piecemeal = reopen(two_edits)
    for rev_id in reversed(order) if reverse else order:
        piecemeal.reject(ids=rev_id)

    assert texts(piecemeal.document) == original_text
    assert len(piecemeal.summary()) == 0


def test_every_survivor_still_resolves_both_ways(two_edits, original_text):
    """Whatever is left after a partial accept must still be a valid redline."""
    for rev_id in two_edits.summary().ids:
        partial = reopen(two_edits)
        partial.accept(ids=rev_id)

        accepted, rejected = reopen(partial), reopen(partial)
        accepted.accept_all()
        rejected.reject_all()

        whole = reopen(two_edits)
        whole.accept_all()
        assert texts(accepted.document) == texts(whole.document), rev_id
        assert len(accepted.summary()) == len(rejected.summary()) == 0


def test_no_filter_is_accept_all(two_edits):
    filtered = reopen(two_edits)
    filtered.accept()
    everything = reopen(two_edits)
    everything.accept_all()
    assert texts(filtered.document) == texts(everything.document)
    assert make_selector() is None


# ---------------------------------------------------------------------------
# selecting by author / kind / predicate
# ---------------------------------------------------------------------------
@pytest.fixture
def two_authors(contract_bytes):
    first = Redliner(contract_bytes, author="Tester", date=DATE)
    first.replace_text("thirty (30) days", "forty-five (45) days")
    second = Redliner(roundtrip(first), author="Opposing Counsel", date=DATE)
    second.replace_text("1.5%", "2.0%")
    return second


def test_accept_by_author(two_authors):
    rl = reopen(two_authors)
    rl.accept(authors="Tester")
    assert set(rl.summary().authors) == {"Opposing Counsel"}
    assert "forty-five (45) days" in "\n".join(texts(rl.document))


def test_reject_by_author(two_authors):
    rl = reopen(two_authors)
    rl.reject(authors=["Opposing Counsel"])
    assert set(rl.summary().authors) == {"Tester"}
    assert "1.5%" in "\n".join(texts(rl.document))


def test_author_and_id_combine_with_and(two_authors):
    rev = next(r for r in two_authors.summary().revisions if r.author == "Tester")
    rl = reopen(two_authors)
    rl.accept(ids=rev.id, authors="Opposing Counsel")  # no revision is both
    assert len(rl.summary()) == len(two_authors.summary())


def test_accept_by_kind_family(rl):
    rl.format_matching("Late Payment", bold=True)
    rl.replace_text("thirty (30) days", "forty-five (45) days")
    other = reopen(rl)
    other.accept(kinds="format")
    assert other.summary().counts["format:rPrChange"] == 0
    assert other.summary().counts["insert"] == 1


def test_reject_only_deletions(two_edits):
    rl = reopen(two_edits)
    rl.reject(kinds="delete")
    assert "thirty (30) days" in "\n".join(texts(rl.document))  # the strikeout is undone
    assert rl.summary().counts["delete"] == 0
    assert rl.summary().counts["insert"] == 2  # the replacement text is still proposed


def test_where_predicate(two_edits):
    rl = reopen(two_edits)
    rl.accept(where=lambda rev: "Notices" in rev.text)
    assert "5. Notices." in "\n".join(texts(rl.document))
    assert rl.summary().counts["insert"] == 1  # the replacement is untouched


def test_kind_alias_covers_paragraph_marks(two_edits):
    rl = reopen(two_edits)
    rl.accept(kinds="insert-any")
    assert rl.summary().counts["insert"] == 0
    assert rl.summary().counts["paragraph-mark-insert"] == 0
    assert rl.summary().counts["delete"] == 1


# ---------------------------------------------------------------------------
# moves must resolve as a pair
# ---------------------------------------------------------------------------
@pytest.fixture
def moved(contract_bytes):
    rl = Redliner(contract_bytes, author="Tester", date=DATE)
    rl.move_paragraph(
        rl.find_paragraph(contains="2. Invoicing"),
        after=rl.find_paragraph(contains="4. Termination"),
    )
    rl.replace_text("1.5%", "2.0%")
    return rl


def test_accepting_one_half_of_a_move_pulls_in_the_other(moved):
    move_from = next(r for r in moved.summary().revisions if r.kind == "move-from" and r.text)
    rl = reopen(moved)
    rl.accept(ids=move_from.id)

    body = texts(rl.document)
    assert sum("2. Invoicing" in line for line in body) == 1  # not duplicated, not lost
    assert position(body, "2. Invoicing") > position(body, "4. Termination")
    assert not [r for r in rl.summary().revisions if r.kind.startswith("move-")]
    assert rl.summary().counts["delete"] == 1  # the unrelated edit survives


def test_rejecting_one_half_of_a_move_pulls_in_the_other(moved, original_text):
    move_to = next(r for r in moved.summary().revisions if r.kind == "move-to" and r.text)
    rl = reopen(moved)
    rl.reject(ids=move_to.id)

    body = texts(rl.document)
    assert sum("2. Invoicing" in line for line in body) == 1
    assert position(body, "2. Invoicing") < position(body, "4. Termination")
    assert not [r for r in rl.summary().revisions if r.kind.startswith("move-")]


def test_resolved_move_drops_its_range_markers(moved):
    rl = reopen(moved)
    rl.accept(kinds="move")
    body = rl.document.element.body
    assert not body.findall(".//" + qn("w:moveFromRangeStart"))
    assert not body.findall(".//" + qn("w:moveToRangeEnd"))


def test_unselected_move_keeps_its_range_markers(moved):
    rl = reopen(moved)
    rl.accept(kinds="delete")  # nothing to do with the move
    body = rl.document.element.body
    assert body.findall(".//" + qn("w:moveFromRangeStart"))
    assert body.findall(".//" + qn("w:moveToRangeStart"))
    assert len([r for r in rl.summary().revisions if r.kind.startswith("move-")]) == 4


def test_selecting_a_range_marker_id_resolves_the_move(moved):
    rl = reopen(moved)
    start = rl.document.element.body.find(".//" + qn("w:moveFromRangeStart"))
    rl.accept(ids=start.get(qn("w:id")))
    assert not [r for r in rl.summary().revisions if r.kind.startswith("move-")]
    assert sum("2. Invoicing" in line for line in texts(rl.document)) == 1


# ---------------------------------------------------------------------------
# table rows and paragraph marks
# ---------------------------------------------------------------------------
def test_selective_row_revisions(rl):
    table = rl.tables()[0]
    rl.insert_table_row(table, values=("Enterprise", "500", "$150,000"))
    rl.delete_table_row(rl.tables()[0], 1)

    inserted = next(r for r in rl.summary().revisions if r.kind == "row-insert")
    other = reopen(rl)
    other.accept(ids=inserted.id)

    assert other.summary().counts["row-insert"] == 0
    assert other.summary().counts["row-delete"] == 1  # still awaiting review
    assert "Enterprise" in "\n".join(texts(other.document))
    assert len(other.tables()[0].rows) == 4  # the struck row is still there to review

    restored = reopen(other)
    restored.reject_all()
    assert "Standard" in "\n".join(texts(restored.document))


def test_selective_paragraph_deletion(rl, original_text):
    rl.delete_paragraph(rl.find_paragraph(contains="3. Late Payment"))
    rl.replace_text("thirty (30) days", "forty-five (45) days")

    mark = next(r for r in rl.summary().revisions if r.kind == "paragraph-mark-delete")
    other = reopen(rl)
    other.accept(kinds=["delete", "paragraph-mark-delete"])

    assert "3. Late Payment" not in "\n".join(texts(other.document))
    assert not other.summary().by_id(mark.id)
    assert other.summary().counts["insert"] == 1  # the replacement is still tracked


def test_rejecting_an_added_paragraph_leaves_no_husk(rl, original_text):
    """The ¶ mark and the runs are separate ids; either order must clean up.

    The inserted paragraph sits right before the table, so there is no following
    paragraph to merge into -- the case that used to strand an empty one.
    """
    rl.insert_paragraph_after(rl.find_paragraph(contains="4. Termination"), "5. Notices.")
    mark = next(r for r in rl.summary().revisions if r.kind == "paragraph-mark-insert")
    runs = next(r for r in rl.summary().revisions if r.kind == "insert")

    other = reopen(rl)
    other.reject(ids=mark.id)  # the mark first -- the awkward order
    other.reject(ids=runs.id)
    assert texts(other.document) == original_text


def test_unselected_paragraph_mark_is_left_alone(rl):
    rl.delete_paragraph(rl.find_paragraph(contains="3. Late Payment"))
    before = rl.summary().counts["paragraph-mark-delete"]
    other = reopen(rl)
    other.accept(ids="999999")
    assert other.summary().counts["paragraph-mark-delete"] == before


# ---------------------------------------------------------------------------
# the low-level and file-level entry points
# ---------------------------------------------------------------------------
def test_module_level_accept_takes_a_raw_predicate(two_edits):
    rl = reopen(two_edits)
    body = rl.document.element.body
    review.accept(body, lambda el: el.get(qn("w:id")) == two_edits.summary().ids[0])
    assert len(rl.summary()) == len(two_edits.summary()) - 1


def test_accept_file_filters(tmp_path, two_edits):
    source = tmp_path / "redlined.docx"
    two_edits.save(source)
    rev = next(r for r in two_edits.summary().revisions if r.text == "5. Notices.")

    out = accept_file(source, tmp_path / "partial.docx", ids=rev.id)
    assert len(Redliner(out, track_changes=False).summary()) == len(two_edits.summary()) - 1
    assert "5. Notices." in "\n".join(texts(docx.Document(str(out))))

    out = reject_file(source, tmp_path / "none.docx", authors="Nobody")
    assert len(Redliner(out, track_changes=False).summary()) == len(two_edits.summary())


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def test_cli_accepts_a_single_id(tmp_path, two_edits, capsys):
    from docx_redline.cli import main

    source = tmp_path / "redlined.docx"
    two_edits.save(source)
    rev = next(r for r in two_edits.summary().revisions if r.text == "5. Notices.")
    out = tmp_path / "partial.docx"

    assert main(["accept", str(source), "-o", str(out), "--id", rev.id]) == 0
    assert "1 of 4 change(s) accepted" in capsys.readouterr().out
    assert "5. Notices." in "\n".join(texts(docx.Document(str(out))))


def test_cli_splits_comma_separated_authors(tmp_path, two_authors, capsys):
    from docx_redline.cli import main

    source = tmp_path / "redlined.docx"
    two_authors.save(source)
    out = tmp_path / "partial.docx"

    assert main(["accept", str(source), "-o", str(out), "--author", "Tester,Nobody"]) == 0
    assert set(Redliner(out, track_changes=False).summary().authors) == {"Opposing Counsel"}


def test_cli_summary_json_carries_ids(tmp_path, two_edits, capsys):
    import json

    from docx_redline.cli import main

    source = tmp_path / "redlined.docx"
    two_edits.save(source)
    assert main(["summary", str(source), "--json"]) == 0
    rows = json.loads(capsys.readouterr().out)
    assert all(row["id"] for row in rows)
